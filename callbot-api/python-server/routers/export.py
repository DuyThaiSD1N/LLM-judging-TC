"""
Export Excel router - Xuất testcases ra Excel
"""

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO
from datetime import datetime

from models.schemas import ExportRequest


router = APIRouter()


@router.post("/export")
async def export_testcases(request: ExportRequest):
    """
    Xuất testcases + kết quả ra Excel
    
    POST /api/export
    Body: { testcases: [...] }
    """
    if not request.testcases:
        raise HTTPException(status_code=400, detail="Không có testcase nào để xuất")
    
    try:
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Testcases"
        
        # Headers - SAME ORDER AS TEMPLATE + Run Time
        headers = [
            "Tên Testcase", "Mã TC", "Người thực hiện", "Setup lịch sử", "Câu hỏi từ User", "Yêu cầu kỳ vọng",
            "Từ khóa BẮT BUỘC", "Từ khóa CẤM", "LLM Judge", "Bot URL",
            "Lượt", "Thời gian chạy", "Câu trả lời thực tế", "Thời gian (ms)", "Kết quả",
            "Lỗi", "Đề xuất sửa", "Mẫu đề xuất", "Nhận xét giọng điệu"
        ]
        ws.append(headers)
        
        # Style header
        header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        header_font = Font(bold=True, size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Criteria names mapping
        criteria_names = {
            "goal_achievement": "Đạt mục tiêu",
            "semantic_correctness": "Đúng nghĩa & intent",
            "conversation_quality": "Chất lượng hội thoại",
            "context_consistency": "Tính nhất quán",
            "safety_compliance": "An toàn & Tuân thủ"
        }
        
        # Data rows - xử lý dict thay vì object
        for tc in request.testcases:
            criteria_display = criteria_names.get(tc.get('criteria', 'standard'), tc.get('criteria', 'standard'))
            bot_url_display = tc.get('bot_url') if tc.get('bot_url') else "mặc định"
            run_at = tc.get('run_at', '')  # Get run time if available
            
            # Format run_at to readable format
            if run_at:
                try:
                    from datetime import datetime
                    dt = datetime.fromisoformat(run_at.replace(' ', 'T')) if 'T' not in run_at else datetime.fromisoformat(run_at)
                    run_at_display = dt.strftime("%d/%m/%Y %H:%M")
                except:
                    run_at_display = run_at
            else:
                run_at_display = ""
            
            turns = tc.get('turns', [])
            for idx, turn in enumerate(turns):
                # Lấy tất cả fields, hỗ trợ cả Turn và TurnResult
                scenario = turn.get('scenario', '') or ''
                required_kw = turn.get('required_keywords', '') or ''
                forbidden_kw = turn.get('forbidden_keywords', '') or ''
                actual = turn.get('actual', '') or ''
                response_time = turn.get('response_time_ms', '') or ''
                verdict = turn.get('verdict', '') or ''
                error_desc = turn.get('error_desc', '') or ''
                suggestion = turn.get('suggestion', '') or ''
                suggested_response = turn.get('suggested_response', '') or ''
                tone_note = turn.get('tone_note', '') or ''
                
                row = [
                    tc.get('name', ''),
                    tc.get('code', ''),
                    tc.get('executor', '') or '',
                    scenario,
                    turn.get('question', ''),
                    turn.get('expected', ''),
                    required_kw,
                    forbidden_kw,
                    criteria_display,
                    bot_url_display,
                    f"Lượt {idx + 1}",
                    run_at_display,  # Run time column
                    actual,
                    response_time,
                    verdict,
                    error_desc,
                    suggestion,
                    suggested_response,
                    tone_note
                ]
                ws.append(row)
        
        # Column widths - 19 cột (thêm Người thực hiện)
        # 1: Tên TC, 2: Mã TC, 3: Người thực hiện, 4: Setup lịch sử, 5: Câu hỏi, 6: Yêu cầu kỳ vọng,
        # 7: Từ khóa BẮT BUỘC, 8: Từ khóa CẤM, 9: LLM Judge, 10: Bot URL, 11: Lượt, 12: Thời gian chạy,
        # 13: Câu trả lời thực tế, 14: Thời gian (ms), 15: Kết quả, 16: Lỗi, 17: Đề xuất sửa, 18: Mẫu đề xuất, 19: Nhận xét giọng điệu
        column_widths = [35, 12, 22, 50, 45, 65, 30, 30, 15, 35, 10, 18, 55, 12, 10, 45, 55, 55, 55]
        for idx, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + idx) if idx <= 26 else f"A{chr(64 + idx - 26)}"].width = width
        
        # Style all data cells: wrap text and top alignment
        data_alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = data_alignment
                cell.border = thin_border

        # Style verdict cells (PASSED = green, FAILED = red) - Column O (15) - shifted +1 for executor
        passed_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        passed_font = Font(color="006100", bold=True)
        failed_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        failed_font = Font(color="9C0006", bold=True)
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=15, max_col=15):
            cell = row[0]
            if cell.value == "PASSED":
                cell.fill = passed_fill
                cell.font = passed_font
            elif cell.value == "FAILED":
                cell.fill = failed_fill
                cell.font = failed_font
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Generate filename with timestamp
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        filename = f"testcases_{timestamp}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

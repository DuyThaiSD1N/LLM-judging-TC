"""
Template router - Xuất file Excel mẫu
"""

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from io import BytesIO


router = APIRouter()


@router.get("/template")
async def get_template():
    """
    Xuất file Excel mẫu để nhập testcase
    
    GET /api/template
    """
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Testcases"
    
    # Headers
    headers = [
        "Tên Testcase",
        "Mã Testcase",
        "Người thực hiện (tùy chọn)",
        "Setup lịch sử (không bắt buộc)",
        "Câu hỏi từ User",
        "Yêu cầu kỳ vọng",
        "Từ khóa BẮT BUỘC (tùy chọn)",
        "Từ khóa CẤM (tùy chọn)",
        "LLM Judge (goal_achievement/semantic_correctness/conversation_quality/context_consistency/safety_compliance)",
        "Bot URL (tùy chọn)"
    ]
    ws.append(headers)
    
    # Sample data
    sample = [
        "Hỏi thủ tục đăng ký kết hôn",
        "TC-001",
        "Nguyễn Văn A",  # Người thực hiện
        "[user] đăng ký lại khai sinh thì làm sao vậy?\n[assistant] Dạ về thủ tục Đăng ký lại khai sinh, người dân nộp hồ sơ tại Ủy ban nhân dân cấp xã hoặc nộp trực tuyến qua Cổng dịch vụ công. Thời hạn giải quyết là 5 Ngày, trường hợp phải xác minh thì không quá 25 ngày ạ.",  # Setup lịch sử - [assistant] chỉ để tham khảo
        "đăng ký kết hôn cần giấy tờ gì",
        "Cần CMND/CCCD và giấy xác nhận tình trạng hôn nhân, nộp tại UBND cấp xã",
        "CMND, giấy xác nhận",  # Từ khóa bắt buộc
        "không rõ, không biết",  # Từ khóa cấm
        "goal_achievement",
        ""  # Bot URL để trống = dùng mặc định
    ]
    ws.append(sample)
    
    # Column widths (10 cột: Tên, Mã, Người thực hiện, Setup, Câu hỏi, Kỳ vọng, KW bắt buộc, KW cấm, Judge, Bot URL)
    column_widths = [35, 12, 22, 50, 45, 60, 30, 30, 50, 40]
    for idx, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    
    # Style header
    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=testcase_mau.xlsx"}
    )
